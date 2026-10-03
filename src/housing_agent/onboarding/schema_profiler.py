from __future__ import annotations

import csv
import json
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable

from openpyxl import load_workbook

from housing_agent.onboarding.mapping_models import (
    FieldProfile,
    RawSchemaProfile,
    TableProfile,
)


PERIOD_FIELD = re.compile(
    r"(?:\d{4}(?:[-/]?(?:0?[1-9]|1[0-2])|[- ]?Q[1-4]|(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)))",
    re.IGNORECASE,
)


def _json_safe(value: Any) -> Any:
    if isinstance(value, (date, datetime)):
        return value.isoformat()
    if isinstance(value, (str, int, float, bool)) or value is None:
        return value
    return str(value)


def _flatten(record: dict[str, Any], prefix: str = "") -> dict[str, Any]:
    output: dict[str, Any] = {}
    for key, value in record.items():
        path = f"{prefix}.{key}" if prefix else str(key)
        if isinstance(value, dict):
            output.update(_flatten(value, path))
        elif not isinstance(value, (list, tuple, set)):
            output[path] = value
    return output


def _type_name(value: Any) -> str:
    if value is None:
        return "null"
    if isinstance(value, bool):
        return "boolean"
    if isinstance(value, int):
        return "integer"
    if isinstance(value, float):
        return "number"
    if isinstance(value, (date, datetime)):
        return "date"
    return "string"


def _field_profiles(records: list[dict[str, Any]]) -> list[FieldProfile]:
    flattened = [_flatten(record) for record in records]
    fields = sorted({key for record in flattened for key in record})
    output: list[FieldProfile] = []
    for field in fields:
        values = [record.get(field) for record in flattened]
        examples: list[str] = []
        for value in values:
            if value is None:
                continue
            text = str(_json_safe(value))[:120]
            if text not in examples:
                examples.append(text)
            if len(examples) == 5:
                break
        output.append(FieldProfile(
            path=field,
            observed_types=sorted({_type_name(value) for value in values}),
            non_null_count=sum(value is not None for value in values),
            null_count=sum(value is None for value in values),
            examples=examples,
        ))
    return output


def _period_like_fields(fields: Iterable[FieldProfile]) -> list[str]:
    output: list[str] = []
    for field in fields:
        leaf = field.path.rsplit(".", 1)[-1]
        if PERIOD_FIELD.fullmatch(leaf) or any(
            PERIOD_FIELD.fullmatch(example) for example in field.examples
        ):
            output.append(field.path)
        elif any(token in leaf.lower() for token in ("date", "period", "month", "quarter", "year")):
            output.append(field.path)
    return sorted(output)


def _table_profile(
    *,
    location: str,
    records: list[dict[str, Any]],
    records_path: tuple[str | int, ...] = (),
    sheet_name: str | None = None,
    header_row: int | None = None,
) -> TableProfile:
    profiles = _field_profiles(records)
    period_fields = _period_like_fields(profiles)
    period_column_names = [item for item in period_fields if PERIOD_FIELD.fullmatch(item)]
    likely = "wide" if len(period_column_names) >= 2 else "long" if period_fields else "unknown"
    return TableProfile(
        location=location,
        records_path=records_path,
        sheet_name=sheet_name,
        header_row=header_row,
        record_count=len(records),
        fields=profiles,
        period_like_fields=period_fields,
        likely_layout=likely,
        sample_records=[
            {key: _json_safe(value) for key, value in _flatten(record).items()}
            for record in records[:5]
        ],
    )


def _json_tables(payload: Any) -> list[TableProfile]:
    found: list[TableProfile] = []

    def visit(value: Any, path: tuple[str | int, ...]) -> None:
        if isinstance(value, list):
            records = [item for item in value if isinstance(item, dict)]
            if records:
                location = "$" + "".join(
                    f"[{part}]" if isinstance(part, int) else f".{part}" for part in path
                )
                found.append(_table_profile(
                    location=location,
                    records=records,
                    records_path=path,
                ))
                return
            for index, child in enumerate(value[:20]):
                visit(child, (*path, index))
        elif isinstance(value, dict):
            for key, child in value.items():
                visit(child, (*path, str(key)))

    visit(payload, ())
    if not found and isinstance(payload, dict):
        found.append(_table_profile(location="$", records=[payload]))
    return found


def _detect_excel_header(rows: list[tuple[Any, ...]]) -> int:
    candidates = rows[:20]
    scores: list[tuple[int, int]] = []
    for index, row in enumerate(candidates, start=1):
        non_empty = [value for value in row if value not in (None, "")]
        text_count = sum(isinstance(value, str) for value in non_empty)
        unique_count = len({str(value) for value in non_empty})
        scores.append((text_count * 3 + unique_count, index))
    return max(scores)[1] if scores else 1


def profile_snapshot(
    *, source_id: str, snapshot_path: str | Path, snapshot_sha256: str
) -> RawSchemaProfile:
    path = Path(snapshot_path)
    suffix = path.suffix.lower()
    warnings: list[str] = []
    tables: list[TableProfile] = []
    if suffix == ".json":
        source_format = "json"
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
        tables = _json_tables(payload)
    elif suffix == ".csv":
        source_format = "csv"
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            records = list(reader)
        if not reader.fieldnames or not records:
            raise ValueError("CSV snapshot has no header or data records")
        tables = [_table_profile(location="csv:root", records=records, header_row=1)]
    elif suffix in {".xlsx", ".xlsm"}:
        source_format = "xlsx"
        workbook = load_workbook(path, read_only=True, data_only=True)
        for sheet in workbook.worksheets:
            values = list(sheet.iter_rows(values_only=True))
            if not any(any(value not in (None, "") for value in row) for row in values):
                warnings.append(f"Skipped empty sheet {sheet.title!r}.")
                continue
            header_row = _detect_excel_header(values)
            headers = [
                str(value).strip() if value not in (None, "") else f"__column_{index + 1}"
                for index, value in enumerate(values[header_row - 1])
            ]
            records = [
                dict(zip(headers, row, strict=False))
                for row in values[header_row:]
                if any(value not in (None, "") for value in row)
            ]
            if records:
                tables.append(_table_profile(
                    location=f"sheet:{sheet.title}",
                    records=records,
                    sheet_name=sheet.title,
                    header_row=header_row,
                ))
    else:
        raise ValueError(f"Unsupported raw snapshot format: {suffix or '<none>'}")
    if not tables:
        raise ValueError("Raw snapshot contains no record-like table")
    return RawSchemaProfile(
        source_id=source_id,
        snapshot_path=str(path),
        snapshot_sha256=snapshot_sha256,
        source_format=source_format,
        byte_count=path.stat().st_size,
        tables=tables,
        profiler_warnings=warnings,
    )

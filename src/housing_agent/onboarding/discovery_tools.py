from __future__ import annotations

import csv
import io
import json
import re
import urllib.request
from email.message import Message
from typing import Any, Callable

from openpyxl import load_workbook

from housing_agent.onboarding.discovery_models import (
    DataSourceCandidate,
    SourceProbeResult,
)
from housing_agent.retrieval import USER_AGENT, validate_official_url


def _content_type(headers: Message | Any) -> str | None:
    if hasattr(headers, "get_content_type"):
        return headers.get_content_type()
    value = headers.get("Content-Type") if headers else None
    return value.split(";", 1)[0].strip() if value else None


def _profile_json(body: bytes) -> tuple[list[str], int, int | None, list[dict[str, Any]]]:
    payload = json.loads(body)
    records: object
    total: int | None = None
    if isinstance(payload, dict) and isinstance(payload.get("result"), dict):
        result = payload["result"]
        records = result.get("records", [])
        if result.get("total") is not None:
            total = int(result["total"])
    elif isinstance(payload, list):
        records = payload
    elif isinstance(payload, dict):
        records = [payload]
    else:
        raise ValueError("JSON response is not an object or array")
    if not isinstance(records, list) or not records:
        raise ValueError("JSON response contains no sample records")
    dictionaries = [item for item in records if isinstance(item, dict)]
    if not dictionaries:
        raise ValueError("JSON sample records are not objects")
    fields = sorted({str(key) for item in dictionaries[:25] for key in item})
    return fields, len(dictionaries), total, dictionaries[:25]


def _profile_csv(body: bytes) -> tuple[list[str], int, int | None, list[dict[str, Any]]]:
    text = body.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text))
    rows = list(reader)
    if not reader.fieldnames:
        raise ValueError("CSV response has no header")
    if not rows:
        raise ValueError("CSV response has no sample rows")
    return list(reader.fieldnames), len(rows), None, rows[:25]


def _profile_xlsx(body: bytes) -> tuple[list[str], int, int | None, list[dict[str, Any]]]:
    workbook = load_workbook(io.BytesIO(body), read_only=True, data_only=True)
    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        candidates = []
        for index, row in enumerate(rows[:20]):
            non_empty = [value for value in row if value not in (None, "")]
            score = sum(isinstance(value, str) for value in non_empty) * 3 + len(non_empty)
            candidates.append((score, index))
        if not candidates:
            continue
        _, header_index = max(candidates)
        headers = [
            str(value).strip() if value not in (None, "") else f"__column_{index + 1}"
            for index, value in enumerate(rows[header_index])
        ]
        records = [
            dict(zip(headers, row, strict=False))
            for row in rows[header_index + 1:]
            if any(value not in (None, "") for value in row)
        ]
        if records:
            return headers, len(records), None, records[:25]
    raise ValueError("XLSX response contains no non-empty tabular sheet")


def _normalized_field(value: str) -> str:
    return "".join(character.lower() for character in value if character.isalnum())


MONTH_NUMBER = {
    name: index for index, name in enumerate(
        ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"),
        start=1,
    )
}


def _wide_period_ordinal(value: str, frequency: str | None) -> int:
    if frequency == "singstat_month":
        return int(value[:4]) * 12 + MONTH_NUMBER[value[4:]] - 1
    return int(value[:4]) * 4 + int(value[-1]) - 1


def _has_observed_value(value: object) -> bool:
    return value is not None and str(value).strip().lower() not in {
        "", "na", "n.a.", "null", "-", ".."
    }


def _resolve_proposed_fields(
    proposed: list[str], actual: list[str]
) -> tuple[list[str], list[str]]:
    matches: list[str] = []
    adjustments: list[str] = []
    for name in proposed:
        if name in actual:
            matches.append(name)
            continue
        normalized = _normalized_field(name)
        candidates = [item for item in actual if _normalized_field(item) == normalized]
        if len(candidates) == 1:
            matches.append(candidates[0])
            adjustments.append(f"Resolved proposed field {name!r} to actual field {candidates[0]!r}.")
    return sorted(set(matches)), adjustments


def probe_data_source(
    candidate: DataSourceCandidate,
    *,
    timeout: float = 20.0,
    max_sample_bytes: int = 2_000_000,
    opener: Callable[..., Any] = urllib.request.urlopen,
) -> SourceProbeResult:
    errors: list[str] = []
    warnings: list[str] = []
    final_url: str | None = None
    status: int | None = None
    content_type: str | None = None
    body = b""
    fields: list[str] = []
    sample_count = 0
    total_count: int | None = None
    detected_format: str | None = None
    sample_records: list[dict[str, Any]] = []
    field_adjustments: list[str] = []
    series_field_match: str | None = None
    observed_period_start: str | None = None
    observed_period_end: str | None = None
    non_null_observation_count: int | None = None
    missing_observation_count: int | None = None
    coverage_profile_scope = "unavailable"
    try:
        validate_official_url(candidate.source_page_url)
        validate_official_url(candidate.retrieval_url)
        request = urllib.request.Request(
            candidate.retrieval_url,
            headers={
                "Accept": "application/json,text/csv;q=0.9,*/*;q=0.5",
                "User-Agent": USER_AGENT,
            },
        )
        with opener(request, timeout=timeout) as response:
            status = int(response.status)
            final_url = response.geturl()
            validate_official_url(final_url)
            content_type = _content_type(response.headers)
            body = response.read(max_sample_bytes + 1)
        if len(body) > max_sample_bytes:
            raise ValueError(f"sample response exceeds {max_sample_bytes} bytes")
        if content_type == "application/json" or candidate.source_type.value.endswith("json"):
            detected_format = "json"
            fields, sample_count, total_count, sample_records = _profile_json(body)
        elif content_type in {"text/csv", "application/csv"} or candidate.source_type.value == "csv":
            detected_format = "csv"
            fields, sample_count, total_count, sample_records = _profile_csv(body)
        elif candidate.source_type.value == "xlsx" or content_type in {
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/vnd.ms-excel",
        }:
            detected_format = "xlsx"
            fields, sample_count, total_count, sample_records = _profile_xlsx(body)
        else:
            detected_format = "unsupported"
            raise ValueError(
                f"discovery probe does not yet profile content type {content_type!r}"
            )
    except Exception as exc:
        errors.append(f"{type(exc).__name__}: {exc}")

    if candidate.schema_orientation == "wide_period_columns":
        period_pattern = (
            r"^\d{4}(?:Jan|Feb|Mar|Apr|May|Jun|Jul|Aug|Sep|Oct|Nov|Dec)$"
            if candidate.wide_period_format == "singstat_month"
            else r"^\d{4}Q[1-4]$"
        )
        period_fields = sorted(field for field in fields if re.fullmatch(period_pattern, field))
        date_matches = period_fields
        value_matches = period_fields
        if candidate.series_field:
            series_matches, series_adjustments = _resolve_proposed_fields(
                [candidate.series_field], fields
            )
            field_adjustments.extend(series_adjustments)
            if series_matches:
                series_field_match = series_matches[0]
            elif not errors:
                warnings.append(
                    f"Proposed series field {candidate.series_field!r} is missing; full-snapshot mapping is deferred."
                )
            if (
                series_field_match
                and candidate.exact_series_label
                and candidate.exact_series_label not in {
                    str(item.get(series_field_match)) for item in sample_records
                }
                and not errors
            ):
                warnings.append(
                    f"Exact series label {candidate.exact_series_label!r} was not found in the sample; full-snapshot mapping is deferred."
                )
            elif series_field_match and candidate.exact_series_label:
                target_rows = [
                    item for item in sample_records
                    if str(item.get(series_field_match)) == candidate.exact_series_label
                ]
                if len(target_rows) == 1:
                    target = target_rows[0]
                    observed_periods = [
                        field for field in period_fields
                        if _has_observed_value(target.get(field))
                    ]
                    if observed_periods:
                        ordered = sorted(
                            observed_periods,
                            key=lambda value: _wide_period_ordinal(
                                value, candidate.wide_period_format
                            ),
                        )
                        observed_period_start = ordered[0]
                        observed_period_end = ordered[-1]
                        first_ordinal = _wide_period_ordinal(
                            ordered[0], candidate.wide_period_format
                        )
                        last_ordinal = _wide_period_ordinal(
                            ordered[-1], candidate.wide_period_format
                        )
                        missing_observation_count = (
                            last_ordinal - first_ordinal + 1 - len(observed_periods)
                        )
                    non_null_observation_count = len(observed_periods)
                    coverage_profile_scope = "full_response_series"
    else:
        date_matches, date_adjustments = _resolve_proposed_fields(
            candidate.date_field_candidates, fields
        )
        value_matches, value_adjustments = _resolve_proposed_fields(
            candidate.value_field_candidates, fields
        )
        field_adjustments.extend(date_adjustments + value_adjustments)
    if not errors and not date_matches:
        warnings.append(
            "None of the proposed date fields matched the sample; the Mapping Agent must resolve the complete snapshot."
        )
    if not errors and not value_matches:
        warnings.append(
            "None of the proposed value fields matched the sample; the Mapping Agent must resolve the complete snapshot."
        )
    return SourceProbeResult(
        candidate_id=candidate.candidate_id,
        status="fail" if errors else "pass",
        requested_url=candidate.retrieval_url,
        final_url=final_url,
        http_status=status,
        content_type=content_type,
        bytes_sampled=len(body),
        detected_format=detected_format,
        fields=fields,
        sample_record_count=sample_count,
        total_record_count=total_count,
        date_field_matches=date_matches,
        value_field_matches=value_matches,
        series_field_match=series_field_match,
        field_adjustments=field_adjustments,
        observed_period_start=observed_period_start,
        observed_period_end=observed_period_end,
        non_null_observation_count=non_null_observation_count,
        missing_observation_count=missing_observation_count,
        coverage_profile_scope=coverage_profile_scope,
        errors=errors,
        warnings=warnings,
    )

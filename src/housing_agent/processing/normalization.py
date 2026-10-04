from __future__ import annotations

import csv
import json
import re
from calendar import monthrange
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable

from housing_agent.processing.models import (
    CanonicalObservation,
    FrozenDataPlan,
    FrozenSeriesPlan,
    MappingContract,
    TransformationTrace,
)
from housing_agent.processing.policy_registry import DerivedOperation


class MappingFailure(ValueError):
    """The approved mapping no longer describes the raw snapshot."""


MONTHS = {
    name.lower(): number
    for number, name in enumerate(
        ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec"),
        start=1,
    )
}


def _json_value(payload: Any, path: Iterable[str | int]) -> Any:
    current = payload
    for part in path:
        try:
            current = current[part]
        except (KeyError, IndexError, TypeError) as exc:
            raise MappingFailure(f"records_path component {part!r} does not exist") from exc
    return current


def get_record_value(record: dict[str, Any], field_path: str) -> Any:
    """Resolve an exact key first, then a dotted nested-field path."""
    if field_path in record:
        return record[field_path]
    current: Any = record
    for part in field_path.split("."):
        if not isinstance(current, dict) or part not in current:
            return None
        current = current[part]
    return current


def _load_xlsx(path: Path, sheet_name: str | None, header_row: int) -> list[dict[str, Any]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - depends on optional runtime package
        raise RuntimeError("openpyxl is required to normalize xlsx snapshots") from exc
    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet_name and sheet_name not in workbook.sheetnames:
        raise MappingFailure(f"approved Excel sheet {sheet_name!r} no longer exists")
    sheet = workbook[sheet_name] if sheet_name else workbook[workbook.sheetnames[0]]
    rows = sheet.iter_rows(values_only=True)
    for _ in range(header_row - 1):
        next(rows, None)
    headers = next(rows, None)
    if not headers:
        raise MappingFailure("Excel snapshot has no approved header row")
    names = [str(value).strip() if value is not None else "" for value in headers]
    return [dict(zip(names, row, strict=False)) for row in rows]


def load_mapping_records(
    snapshot_path: str | Path, contract: Any
) -> list[dict[str, Any]]:
    """Load the record set selected by an approved or proposed mapping contract."""
    path = Path(snapshot_path)
    if contract.source_format == "json":
        payload = json.loads(path.read_text(encoding="utf-8-sig"))
        selected = _json_value(payload, contract.records_path)
        if isinstance(selected, dict):
            records = [selected]
        elif isinstance(selected, list):
            records = selected
        else:
            raise MappingFailure("approved records_path does not resolve to an object or list")
    elif contract.source_format == "csv":
        with path.open("r", encoding="utf-8-sig", newline="") as handle:
            records = list(csv.DictReader(handle))
    else:
        records = _load_xlsx(path, contract.sheet_name, contract.header_row)
    records = [item for item in records if isinstance(item, dict)]
    if contract.series_filter_field:
        if not any(
            get_record_value(item, contract.series_filter_field) is not None
            for item in records
        ):
            raise MappingFailure(
                f"approved series filter field {contract.series_filter_field!r} is missing"
            )
        records = [
            item for item in records
            if str(get_record_value(item, contract.series_filter_field))
            == contract.series_filter_value
        ]
    for field, expected in contract.additional_series_filters.items():
        if not any(get_record_value(item, field) is not None for item in records):
            raise MappingFailure(f"approved additional series filter field {field!r} is missing")
        records = [
            item for item in records
            if str(get_record_value(item, field)) == expected
        ]
    if not records:
        raise MappingFailure("approved mapping selected no records")
    return records


def load_raw_records(series: FrozenSeriesPlan) -> list[dict[str, Any]]:
    return load_mapping_records(series.snapshot_path, series.mapping_contract)


def parse_decimal(value: Any, missing_markers: tuple[str, ...]) -> Decimal | None:
    if value is None:
        return None
    text = str(value).strip()
    if text.lower() in {marker.strip().lower() for marker in missing_markers}:
        return None
    negative = text.startswith("(") and text.endswith(")")
    cleaned = text.strip("()").replace(",", "").replace("%", "").strip()
    try:
        parsed = Decimal(cleaned)
    except InvalidOperation as exc:
        raise MappingFailure(f"value {text!r} cannot be parsed as a decimal") from exc
    return -parsed if negative else parsed


def parse_period(value: Any, frequency: str) -> tuple[str, date]:
    text = str(value).strip()
    if frequency == "annual":
        match = re.fullmatch(r"(\d{4})", text)
        if match:
            year = int(match.group(1))
            return str(year), date(year, 12, 31)
    elif frequency == "quarterly":
        match = re.fullmatch(r"(\d{4})[- ]?Q([1-4])", text, re.IGNORECASE)
        singstat_match = re.fullmatch(r"(\d{4})([1-4])Q", text, re.IGNORECASE)
        if match:
            year, quarter = int(match.group(1)), int(match.group(2))
            month = quarter * 3
            return f"{year}-Q{quarter}", date(year, month, monthrange(year, month)[1])
        if singstat_match:
            year, quarter = int(singstat_match.group(1)), int(singstat_match.group(2))
            month = quarter * 3
            return f"{year}-Q{quarter}", date(year, month, monthrange(year, month)[1])
    elif frequency == "monthly":
        numeric = re.fullmatch(r"(\d{4})[-/]?(0?[1-9]|1[0-2])", text)
        named = re.fullmatch(r"(\d{4})[- ]?([A-Za-z]{3})", text)
        if numeric:
            year, month = int(numeric.group(1)), int(numeric.group(2))
            return f"{year}-{month:02d}", date(year, month, monthrange(year, month)[1])
        if named and named.group(2).lower() in MONTHS:
            year, month = int(named.group(1)), MONTHS[named.group(2).lower()]
            return f"{year}-{month:02d}", date(year, month, monthrange(year, month)[1])
    raise MappingFailure(f"period {text!r} does not match {frequency} frequency")


def _parse_release_date(value: Any) -> date | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value).strip()[:10])
    except ValueError as exc:
        raise MappingFailure(f"release date {value!r} is not ISO YYYY-MM-DD") from exc


def _eligibility(
    period_end: date,
    value: Decimal | None,
    available_date: date,
    plan: FrozenDataPlan,
) -> str:
    if period_end < plan.historical_start_date:
        return "before_window"
    if value is None:
        return "missing_at_source"
    if period_end > plan.historical_end_date or available_date > plan.reporting_as_of_date:
        return "after_as_of"
    return "eligible"


def _emit(
    *,
    plan: FrozenDataPlan,
    series: FrozenSeriesPlan,
    period_value: Any,
    raw_value: Any,
    raw_reference: str,
    release_value: Any = None,
    contract: MappingContract | None = None,
) -> CanonicalObservation:
    active_contract = contract or series.mapping_contract
    period, period_end = parse_period(period_value, series.observation_frequency)
    parsed = parse_decimal(raw_value, active_contract.missing_markers)
    scaled = parsed * active_contract.unit_scale if parsed is not None else None
    release_date = _parse_release_date(release_value)
    available_date = release_date or period_end + timedelta(days=series.release_lag_days)
    return CanonicalObservation(
        source_id=series.source_id,
        series_id=series.series_id,
        series_label=series.series_label,
        period=period,
        value=scaled,
        unit=series.unit,
        observation_frequency=series.observation_frequency,
        update_frequency=series.update_frequency,
        portfolio_role=series.portfolio_role,
        market_role=series.market_role,
        eligibility_status=_eligibility(period_end, scaled, available_date, plan),
        raw_snapshot_sha256=series.snapshot_sha256,
        raw_record_reference=raw_reference,
        original_period=str(period_value),
        original_value=None if raw_value is None else str(raw_value),
        release_date=release_date,
        available_date=available_date,
        source_periods=(period,),
        input_series_ids=(series.series_id,),
    )


def _normalize_with_contract(
    plan: FrozenDataPlan,
    series: FrozenSeriesPlan,
    contract: MappingContract,
    *,
    reference_prefix: str = "",
) -> list[CanonicalObservation]:
    records = load_mapping_records(series.snapshot_path, contract)
    output: list[CanonicalObservation] = []
    if contract.layout == "long":
        assert contract.date_field and contract.value_field
        for field in (contract.date_field, contract.value_field):
            if not any(get_record_value(record, field) is not None for record in records):
                raise MappingFailure(f"approved field {field!r} is missing")
        for index, record in enumerate(records):
            output.append(_emit(
                plan=plan,
                series=series,
                period_value=get_record_value(record, contract.date_field),
                raw_value=get_record_value(record, contract.value_field),
                release_value=(
                    get_record_value(record, contract.release_date_field)
                    if contract.release_date_field else None
                ),
                raw_reference=(
                    f"{reference_prefix}record:{index}/field:{contract.value_field}"
                ),
                contract=contract,
            ))
    else:
        assert contract.wide_period_pattern
        try:
            pattern = re.compile(contract.wide_period_pattern)
        except re.error as exc:
            raise MappingFailure("approved wide_period_pattern is invalid") from exc
        for index, record in enumerate(records):
            period_fields = [field for field in record if pattern.fullmatch(str(field))]
            if not period_fields:
                raise MappingFailure("approved wide mapping found no period fields")
            for field in period_fields:
                output.append(_emit(
                    plan=plan,
                    series=series,
                    period_value=field,
                    raw_value=record.get(field),
                    raw_reference=f"{reference_prefix}record:{index}/field:{field}",
                    contract=contract,
                ))
    return sorted(output, key=lambda item: (item.period, item.raw_record_reference))


def _derived_value(
    operation: DerivedOperation,
    values: dict[str, Decimal],
    scale: Decimal,
) -> Decimal:
    if operation in {DerivedOperation.RATIO, DerivedOperation.SHARE_OF_TOTAL}:
        denominator = values["denominator"]
        if denominator == 0:
            raise MappingFailure("derived ratio cannot use a zero denominator")
        return values["numerator"] / denominator * scale
    if operation == DerivedOperation.DIFFERENCE:
        return (values["minuend"] - values["subtrahend"]) * scale
    return sum(values.values(), Decimal("0")) * scale


def _combine_eligibility(rows: list[CanonicalObservation]) -> str:
    statuses = {item.eligibility_status for item in rows}
    for status in ("before_window", "after_as_of", "missing_at_source"):
        if status in statuses:
            return status
    return "eligible"


def _normalize_derived_series(
    plan: FrozenDataPlan,
    series: FrozenSeriesPlan,
) -> tuple[list[CanonicalObservation], list[TransformationTrace]]:
    derived = series.derived_series_contract
    assert derived is not None
    inputs: dict[str, dict[str, CanonicalObservation]] = {}
    for input_name, filter_value in derived.input_series_filters.items():
        contract = series.mapping_contract.model_copy(
            update={"series_filter_value": filter_value}
        )
        rows = _normalize_with_contract(
            plan,
            series,
            contract,
            reference_prefix=f"input:{input_name}/",
        )
        inputs[input_name] = {item.period: item for item in rows}

    period_sets = {name: set(rows) for name, rows in inputs.items()}
    all_periods = set().union(*period_sets.values())
    for name, periods in period_sets.items():
        if periods != all_periods:
            missing = sorted(all_periods - periods)
            raise MappingFailure(
                f"derived input {name!r} is missing period(s): {missing[:5]}"
            )

    output: list[CanonicalObservation] = []
    traces: list[TransformationTrace] = []
    for period in sorted(all_periods):
        rows = [inputs[name][period] for name in derived.input_series_filters]
        eligibility = _combine_eligibility(rows)
        values = {
            name: inputs[name][period].value
            for name in derived.input_series_filters
        }
        value = None
        if all(item is not None for item in values.values()):
            value = _derived_value(
                derived.operation,
                {name: item for name, item in values.items() if item is not None},
                derived.scale,
            )
        base = rows[0]
        source_periods = tuple(item.period for item in rows)
        input_ids = tuple(
            f"{series.series_id}:{name}" for name in derived.input_series_filters
        )
        raw_reference = "|".join(item.raw_record_reference for item in rows)
        original_value = ";".join(
            f"{name}={values[name]}" for name in derived.input_series_filters
        )
        output.append(base.model_copy(update={
            "series_id": series.series_id,
            "series_label": series.series_label,
            "value": value,
            "unit": series.unit,
            "eligibility_status": eligibility,
            "raw_record_reference": raw_reference,
            "original_value": original_value,
            "release_date": max(
                (item.release_date for item in rows if item.release_date is not None),
                default=None,
            ),
            "available_date": max(
                (item.available_date for item in rows if item.available_date is not None),
                default=None,
            ),
            "derivation": (
                f"{derived.operation.value} of approved input series; "
                f"scale={derived.scale}"
            ),
            "source_periods": source_periods,
            "input_series_ids": input_ids,
        }))
        traces.append(TransformationTrace(
            operation_id=f"{series.series_id}:derive:{period}",
            operation=f"derive_{derived.operation.value}",
            source_id=series.source_id,
            output_series_id=series.series_id,
            parameters={
                "operation": derived.operation.value,
                "scale": str(derived.scale),
                "input_series_filters": derived.input_series_filters,
                "zero_denominator_policy": derived.zero_denominator_policy,
            },
            input_periods=source_periods,
            output_period=period,
        ))
    return output, traces


def normalize_series_with_trace(
    plan: FrozenDataPlan, series: FrozenSeriesPlan
) -> tuple[list[CanonicalObservation], list[TransformationTrace]]:
    if series.derived_series_contract is not None:
        return _normalize_derived_series(plan, series)
    return _normalize_with_contract(plan, series, series.mapping_contract), []


def normalize_series(plan: FrozenDataPlan, series: FrozenSeriesPlan) -> list[CanonicalObservation]:
    observations, _ = normalize_series_with_trace(plan, series)
    return observations


def normalize_plan_with_trace(
    plan: FrozenDataPlan,
) -> tuple[list[CanonicalObservation], list[TransformationTrace]]:
    output: list[CanonicalObservation] = []
    traces: list[TransformationTrace] = []
    for series in sorted(plan.series, key=lambda item: (item.source_id, item.series_id)):
        observations, series_traces = normalize_series_with_trace(plan, series)
        output.extend(observations)
        traces.extend(series_traces)
    return output, traces


def normalize_plan(plan: FrozenDataPlan) -> list[CanonicalObservation]:
    observations, _ = normalize_plan_with_trace(plan)
    return observations

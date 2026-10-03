from __future__ import annotations

import hashlib
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from housing_agent.onboarding.schema_profiler import profile_snapshot


def _sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


class SchemaProfilerTests(unittest.TestCase):
    def test_nested_json_finds_record_path_and_nested_fields(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source.json"
            path.write_text(
                '{"result":{"records":[{"meta":{"period":"2026-01"},"value":"1.2"},'
                '{"meta":{"period":"2026-02"},"value":"1.3"}]}}',
                encoding="utf-8",
            )

            profile = profile_snapshot(
                source_id="source_a", snapshot_path=path, snapshot_sha256=_sha(path)
            )

            self.assertEqual(profile.source_format, "json")
            self.assertEqual(profile.tables[0].records_path, ("result", "records"))
            self.assertIn("meta.period", {item.path for item in profile.tables[0].fields})
            self.assertEqual(profile.tables[0].likely_layout, "long")

    def test_wide_csv_is_identified_without_source_specific_field_names(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source.csv"
            path.write_text(
                "Series,2026Jan,2026Feb,2026Mar\nTarget,1,2,3\n",
                encoding="utf-8",
            )

            profile = profile_snapshot(
                source_id="source_a", snapshot_path=path, snapshot_sha256=_sha(path)
            )

            self.assertEqual(profile.tables[0].likely_layout, "wide")
            self.assertEqual(
                profile.tables[0].period_like_fields,
                ["2026Feb", "2026Jan", "2026Mar"],
            )

    def test_excel_profiles_each_sheet_and_detects_header_row(self) -> None:
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "source.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Monthly Data"
            sheet.append(["Official title"])
            sheet.append([])
            sheet.append(["Month", "Published value"])
            sheet.append(["2026-01", 1.2])
            sheet.append(["2026-02", 1.3])
            workbook.save(path)

            profile = profile_snapshot(
                source_id="source_a", snapshot_path=path, snapshot_sha256=_sha(path)
            )

            table = profile.tables[0]
            self.assertEqual(table.sheet_name, "Monthly Data")
            self.assertEqual(table.header_row, 3)
            self.assertIn("Month", {item.path for item in table.fields})


if __name__ == "__main__":
    unittest.main()
